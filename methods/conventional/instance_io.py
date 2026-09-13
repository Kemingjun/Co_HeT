"""Read and validate the frozen synthetic HRSP workbook format."""

from __future__ import annotations

import ast
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_COLUMNS = [
    "task_index",
    "source_x",
    "source_y",
    "deadline",
    "t_operation",
    "required_robot",
]


@dataclass(frozen=True)
class Task:
    index: int
    x: float
    y: float
    deadline: float
    processing_times: tuple[float, ...]
    required_robot: tuple[int, ...]


@dataclass(frozen=True)
class Instance:
    n: int
    kappa: int
    tasks: tuple[Task, ...]


def _parse_vector(value: Any, label: str, path: Path) -> tuple[Any, ...]:
    try:
        parsed = ast.literal_eval(str(value))
    except (SyntaxError, ValueError) as exc:
        raise ValueError(f"Invalid {label} in {path}: {value!r}") from exc
    if not isinstance(parsed, (list, tuple)):
        raise ValueError(f"Invalid {label} in {path}: expected list")
    return tuple(parsed)


def read_instance(path: Path, *, n: int, kappa: int) -> Instance:
    path = path.resolve()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        if header != EXPECTED_COLUMNS:
            raise ValueError(f"Unexpected columns in {path}: {header}")
        tasks: list[Task] = []
        for row in rows:
            if all(value is None for value in row):
                continue
            index = int(row[0])
            operations = tuple(float(value) for value in _parse_vector(row[4], "t_operation", path))
            required = tuple(int(value) for value in _parse_vector(row[5], "required_robot", path))
            if len(operations) != kappa:
                raise ValueError(f"t_operation length does not match kappa in {path}")
            if required != (1,) * kappa:
                raise ValueError(f"required_robot must be {[1] * kappa} in {path}")
            x, y = float(row[1]), float(row[2])
            if not (0.0 <= x < 1.0 and 0.0 <= y < 1.0):
                raise ValueError(f"Coordinate out of range in {path}")
            tasks.append(Task(index, x, y, float(row[3]), operations, required))
    finally:
        workbook.close()
    if len(tasks) != n:
        raise ValueError(f"Expected {n} tasks in {path}, found {len(tasks)}")
    if [task.index for task in tasks] != list(range(1, n + 1)):
        raise ValueError(f"Non-contiguous task_index in {path}")
    return Instance(n=n, kappa=kappa, tasks=tuple(tasks))


